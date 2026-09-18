import os
import re
import io
import secrets
import string
from datetime import datetime, timedelta

from flask import Flask, render_template, request, jsonify, send_file, redirect, url_for, flash, session

from flask_login import (LoginManager, login_user, logout_user, login_required,
                          current_user)

from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors as rl_colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_RIGHT, TA_CENTER, TA_LEFT

from flask_wtf import CSRFProtect
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address

from models import (db, User, Subscription, Document, LicenseCode, Payment, Message,
                     PaymentRequest, AdminLog, PLAN_PRICES, create_trial_subscription)
from utils import build_sources_context
from translations import get_translations, TRANSLATIONS
from notifications import check_and_send_expiry_reminders, send_email

SUPPORTED_LANGS = ("ar", "fr", "en")


def current_lang():
    return session.get("lang", "fr")


def tr(key, **kwargs):
    """Translate a key for the current session language, with optional .format() args."""
    text = get_translations(current_lang()).get(key, key)
    return text.format(**kwargs) if kwargs else text

# ── APP SETUP ──────────────────────────────────────────────────────────────────
app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "dev-secret-change-in-production")

# DATA_DIR points to a persistent volume in production (e.g. Railway mounts one at /data).
# Falls back to the app's own folder for local development.
DATA_DIR = os.environ.get("DATA_DIR", os.path.dirname(__file__))
os.makedirs(DATA_DIR, exist_ok=True)
app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///" + os.path.join(DATA_DIR, "eduprompt.db")
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["MAX_CONTENT_LENGTH"] = 15 * 1024 * 1024  # 15 MB uploads

CONTACT_EMAIL = os.environ.get("CONTACT_EMAIL", "haithemcomputing@gmail.com")

db.init_app(app)
csrf = CSRFProtect(app)
limiter = Limiter(get_remote_address, app=app, default_limits=[], storage_uri="memory://")

login_manager = LoginManager(app)
login_manager.login_view = "login"


@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))


@app.before_request
def ensure_lang():
    if "lang" not in session:
        best = request.accept_languages.best_match(SUPPORTED_LANGS)
        session["lang"] = best or "fr"
    login_manager.login_message = tr("login_required_message")


@app.context_processor
def inject_i18n():
    lang = current_lang()
    return dict(lang=lang, t=get_translations(lang), is_rtl=(lang == "ar"), contact_email=CONTACT_EMAIL)


@app.context_processor
def inject_notifications():
    """Feeds the navbar notification bell for both teachers and admins."""
    if not current_user.is_authenticated:
        return dict(notif_unread_msgs=0, notif_expiring=0, notif_reminder=False, notif_payments=0)

    if current_user.role == "admin":
        unread_msgs = Message.query.filter_by(sender="teacher", read_by_admin=False).count()
        now = datetime.utcnow()
        soon = now + timedelta(days=5)
        expiring = Subscription.query.filter(
            Subscription.status == "active",
            Subscription.expires_at.isnot(None),
            Subscription.expires_at >= now,
            Subscription.expires_at <= soon,
        ).count()
        pending_payments = PaymentRequest.query.filter_by(status="pending").count()
        return dict(notif_unread_msgs=unread_msgs, notif_expiring=expiring, notif_reminder=False,
                    notif_payments=pending_payments)
    else:
        unread_msgs = Message.query.filter_by(user_id=current_user.id, sender="admin",
                                               read_by_teacher=False).count()
        sub = current_user.subscription
        days_left = sub.days_until_expiry() if sub else None
        reminder = days_left is not None and 0 <= days_left <= 5
        return dict(notif_unread_msgs=unread_msgs, notif_expiring=0, notif_reminder=reminder,
                    notif_payments=0)


def paginate_list(items, page, per_page=15):
    """Manual pagination for plain Python lists (used where results are assembled
    in Python rather than as a single SQLAlchemy Query)."""
    total = len(items)
    total_pages = max(1, (total + per_page - 1) // per_page)
    page = max(1, min(page, total_pages))
    start = (page - 1) * per_page
    return items[start:start + per_page], page, total_pages, total


@app.route("/set-lang/<lang>")
def set_lang(lang):
    if lang in SUPPORTED_LANGS:
        session["lang"] = lang
        if current_user.is_authenticated:
            current_user.preferred_lang = lang
            db.session.commit()
    return redirect(request.referrer or url_for("home"))


def _migrate_sqlite_schema():
    """Add newly-introduced columns to an existing SQLite db (no full migration
    framework needed for this small app). Safe to run on every startup."""
    import sqlite3
    db_path = os.path.join(DATA_DIR, "eduprompt.db")
    if not os.path.exists(db_path):
        return
    conn = sqlite3.connect(db_path)
    cur = conn.cursor()
    try:
        table_additions = {
            "subscription": {
                "docs_used_today": "INTEGER DEFAULT 0",
                "last_reset_date": "DATE",
                "expiry_reminder_sent": "BOOLEAN DEFAULT 0",
            },
            "user": {
                "gemini_api_key": "VARCHAR(255)",
                "preferred_lang": "VARCHAR(5) DEFAULT 'fr'",
                "is_active": "BOOLEAN DEFAULT 1",
            },
        }
        for table, additions in table_additions.items():
            cur.execute(f"PRAGMA table_info({table})")
            existing_cols = {row[1] for row in cur.fetchall()}
            for col, col_type in additions.items():
                if col not in existing_cols:
                    cur.execute(f"ALTER TABLE {table} ADD COLUMN {col} {col_type}")
                    print(f"✅ Migrated: added {table}.{col}")
        conn.commit()
    finally:
        conn.close()


with app.app_context():
    db.create_all()
    _migrate_sqlite_schema()

    # ── One-time admin bootstrap (idempotent — safe on every restart) ──
    _admin_email = os.environ.get("ADMIN_EMAIL")
    _admin_password = os.environ.get("ADMIN_PASSWORD")
    if _admin_email and _admin_password and not User.query.filter_by(email=_admin_email).first():
        _admin = User(
            first_name=os.environ.get("ADMIN_FIRST_NAME", "Admin"),
            last_name=os.environ.get("ADMIN_LAST_NAME", "HaithemEduAI"),
            email=_admin_email,
            role="admin",
        )
        _admin.set_password(_admin_password)
        db.session.add(_admin)
        db.session.flush()
        db.session.add(Subscription(
            user_id=_admin.id, plan="premium", status="active",
            expires_at=datetime.utcnow() + timedelta(days=3650),
            docs_used=0,
        ))
        db.session.commit()
        print(f"✅ Admin account bootstrapped: {_admin_email}")


def log_payment(user_id, plan, amount=None, note=""):
    """Record a manual payment (bank transfer / CCP) for revenue tracking."""
    db.session.add(Payment(
        user_id=user_id,
        plan=plan,
        amount=amount if amount is not None else PLAN_PRICES.get(plan, 0),
        note=note,
    ))


def log_admin_action(action, target="", details=""):
    """Lightweight audit trail — who did what, when."""
    db.session.add(AdminLog(
        admin_id=current_user.id if current_user.is_authenticated else None,
        action=action,
        target=target,
        details=details,
    ))


def apply_payment_approval(teacher, plan, days=365):
    """Activate or renew a teacher's subscription after a payment has been validated.
    Renewing the same plan extends from the later of (now, current expiry);
    switching plan (or reactivating an expired/trial account) starts a fresh period."""
    sub = teacher.subscription
    if not sub:
        sub = Subscription(user_id=teacher.id)
        db.session.add(sub)
        db.session.flush()

    is_renewal = (sub.plan == plan and sub.status == "active" and not sub.is_expired())
    base = sub.expires_at if (is_renewal and sub.expires_at) else datetime.utcnow()

    sub.plan = plan
    sub.status = "active"
    sub.expires_at = base + timedelta(days=days)
    sub.docs_used = 0
    sub.docs_used_today = 0
    sub.last_reset_date = None
    sub.expiry_reminder_sent = False
    return is_renewal


# ── DAILY REMINDER SCHEDULER (5-day heads-up before subscription expiry) ─────
try:
    from apscheduler.schedulers.background import BackgroundScheduler
    _scheduler = BackgroundScheduler(daemon=True)
    _scheduler.add_job(
        lambda: check_and_send_expiry_reminders(app, db, User, Subscription, CONTACT_EMAIL),
        "interval", hours=24, next_run_time=datetime.utcnow(),
        id="expiry_reminders", replace_existing=True,
    )
    _scheduler.start()
    print("✅ Expiry reminder scheduler started")
except Exception as e:
    print(f"⚠️  Could not start reminder scheduler: {e}")


# ── STATIC REFERENCE DATA ────────────────────────────────────────────────────
DOCUMENT_TYPES = {
    "ar": [
        {"id": "lesson_plan",   "label": "مذكرة درس تفصيلية"},
        {"id": "activity",      "label": "ورقة نشاط / تطبيق"},
        {"id": "quiz",          "label": "اختبار QCM مع التصحيح"},
        {"id": "rubric",        "label": "شبكة تقييم Rubric"},
        {"id": "differentiated","label": "دعم متمايز (3 مستويات)"},
        {"id": "summary",       "label": "ملخص / بطاقة مرجعية"},
        {"id": "course_plan",   "label": "تسلسل وحدة تعليمية"},
        {"id": "homework",      "label": "واجب منزلي مُهيكل"},
    ],
    "fr": [
        {"id": "lesson_plan",   "label": "Fiche de cours détaillée"},
        {"id": "activity",      "label": "Feuille d'activité / exercices"},
        {"id": "quiz",          "label": "QCM avec corrigé"},
        {"id": "rubric",        "label": "Grille d'évaluation Rubric"},
        {"id": "differentiated","label": "Supports différenciés (3 niveaux)"},
        {"id": "summary",       "label": "Résumé / fiche mémo"},
        {"id": "course_plan",   "label": "Progression de séquence"},
        {"id": "homework",      "label": "Devoir maison structuré"},
    ],
    "en": [
        {"id": "lesson_plan",   "label": "Detailed lesson plan"},
        {"id": "activity",      "label": "Activity / worksheet"},
        {"id": "quiz",          "label": "MCQ quiz with answer key"},
        {"id": "rubric",        "label": "Assessment rubric"},
        {"id": "differentiated","label": "Differentiated supports (3 levels)"},
        {"id": "summary",       "label": "Summary / reference card"},
        {"id": "course_plan",   "label": "Unit sequence plan"},
        {"id": "homework",      "label": "Structured homework"},
    ],
}

LEVELS = {
    "ar": ["ابتدائي","متوسط","ثانوي","تكوين مهني","جامعي"],
    "fr": ["Primaire","Moyen","Lycée","Formation professionnelle","Université"],
    "en": ["Primary","Middle","High school","Vocational training","University"],
}

SUBJECTS = {
    "ar": ["رياضيات","علوم","تاريخ وجغرافيا","لغة عربية","فيزياء","كيمياء","إعلام آلي","لغة فرنسية","لغة إنجليزية","فلسفة","اقتصاد","أخرى"],
    "fr": ["Mathématiques","Sciences","Histoire-Géographie","Arabe","Physique","Chimie","Informatique","Français","Anglais","Philosophie","Économie","Autre"],
    "en": ["Mathematics","Science","History-Geography","Arabic","Physics","Chemistry","Computer Science","French","English","Philosophy","Economics","Other"],
}


# ── AUTH ROUTES ───────────────────────────────────────────────────────────────
@app.route("/")
def home():
    return render_template("landing.html")


@app.route("/about")
def about():
    from cv_data import CV
    return render_template("about.html", cv=CV.get(current_lang(), CV["fr"]))


@app.route("/register", methods=["GET", "POST"])
@limiter.limit("15/hour")
def register():
    if current_user.is_authenticated:
        return redirect(url_for("generator"))

    if request.method == "POST":
        first_name = request.form.get("first_name", "").strip()
        last_name  = request.form.get("last_name", "").strip()
        email      = request.form.get("email", "").strip().lower()
        password   = request.form.get("password", "")
        api_key    = request.form.get("api_key", "").strip()

        if not all([first_name, last_name, email, password, api_key]):
            flash(tr("flash_all_fields_required"), "error")
            return render_template("register.html")

        if len(password) < 6:
            flash(tr("flash_password_min_length"), "error")
            return render_template("register.html")

        if User.query.filter_by(email=email).first():
            flash(tr("flash_email_exists"), "error")
            return render_template("register.html")

        user = User(first_name=first_name, last_name=last_name, email=email,
                    gemini_api_key=api_key, preferred_lang=current_lang())
        user.set_password(password)
        db.session.add(user)
        db.session.flush()  # get user.id before commit
        create_trial_subscription(user)
        db.session.commit()

        login_user(user)
        flash(tr("flash_welcome_trial"), "success")
        return redirect(url_for("generator"))

    return render_template("register.html")


@app.route("/login", methods=["GET", "POST"])
@limiter.limit("20/hour")
def login():
    if current_user.is_authenticated:
        return redirect(url_for("generator"))

    if request.method == "POST":
        email    = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        user = User.query.filter_by(email=email).first()

        if user and user.check_password(password):
            if not user.is_active:
                flash(tr("flash_account_deactivated"), "error")
                return render_template("login.html")
            login_user(user, remember=True)
            return redirect(url_for("generator"))
        flash(tr("flash_wrong_credentials"), "error")

    return render_template("login.html")


@app.route("/logout")
@login_required
def logout():
    logout_user()
    return redirect(url_for("login"))


# ── DASHBOARD ─────────────────────────────────────────────────────────────────
@app.route("/dashboard")
@login_required
def dashboard():
    sub = current_user.subscription
    recent_docs = current_user.documents[:10]
    recent_messages = Message.query.filter_by(user_id=current_user.id) \
        .order_by(Message.created_at.desc()).limit(3).all()
    unread_admin_msgs = Message.query.filter_by(user_id=current_user.id, sender="admin",
                                                 read_by_teacher=False).count()
    return render_template("dashboard.html", sub=sub, docs=recent_docs,
                            recent_messages=recent_messages, unread_admin_msgs=unread_admin_msgs)


# ── PROFILE (self-service) ───────────────────────────────────────────────────
@app.route("/profile", methods=["GET", "POST"])
@login_required
def profile():
    if request.method == "POST":
        action = request.form.get("action")

        if action == "update_info":
            first_name = request.form.get("first_name", "").strip()
            last_name  = request.form.get("last_name", "").strip()
            email      = request.form.get("email", "").strip().lower()

            if not all([first_name, last_name, email]):
                flash(tr("flash_all_fields_required"), "error")
            else:
                existing = User.query.filter_by(email=email).first()
                if existing and existing.id != current_user.id:
                    flash(tr("flash_email_exists"), "error")
                else:
                    current_user.first_name = first_name
                    current_user.last_name = last_name
                    current_user.email = email
                    db.session.commit()
                    flash(tr("flash_profile_updated"), "success")

        elif action == "change_password":
            current_pw = request.form.get("current_password", "")
            new_pw = request.form.get("new_password", "")
            if not current_user.check_password(current_pw):
                flash(tr("flash_wrong_current_password"), "error")
            elif len(new_pw) < 6:
                flash(tr("flash_password_min_length"), "error")
            else:
                current_user.set_password(new_pw)
                db.session.commit()
                flash(tr("flash_password_changed"), "success")

        elif action == "update_apikey":
            api_key = request.form.get("api_key", "").strip()
            if api_key:
                current_user.gemini_api_key = api_key
                db.session.commit()
                flash(tr("flash_apikey_updated"), "success")

        return redirect(url_for("profile"))

    return render_template("profile.html")


@app.route("/upgrade", methods=["GET", "POST"])
@login_required
def upgrade():
    if request.method == "POST":
        action = request.form.get("action", "code")

        if action == "declare_payment":
            plan = request.form.get("plan", "pro")
            amount = request.form.get("amount", "").strip()
            reference = request.form.get("reference", "").strip()
            note = request.form.get("note", "").strip()

            if plan not in ("pro", "premium") or not reference:
                flash(tr("flash_payment_fields_required"), "error")
                return redirect(url_for("upgrade"))

            db.session.add(PaymentRequest(
                user_id=current_user.id, plan=plan,
                amount_claimed=int(amount) if amount.isdigit() else PLAN_PRICES.get(plan),
                reference=reference, note=note,
            ))
            db.session.commit()
            flash(tr("flash_payment_declared"), "success")
            return redirect(url_for("upgrade"))

        code_str = request.form.get("code", "").strip().upper()
        lic = LicenseCode.query.filter_by(code=code_str, used=False).first()

        if not lic:
            flash(tr("flash_invalid_code"), "error")
            return redirect(url_for("upgrade"))

        apply_payment_approval(current_user, lic.plan, days=lic.duration_days)

        lic.used = True
        lic.used_by = current_user.id
        lic.used_at = datetime.utcnow()

        log_payment(current_user.id, lic.plan, note=f"License code {lic.code}")

        db.session.commit()
        flash(tr("flash_upgraded", plan=current_user.subscription.plan_label()), "success")
        return redirect(url_for("dashboard"))

    my_requests = PaymentRequest.query.filter_by(user_id=current_user.id) \
        .order_by(PaymentRequest.created_at.desc()).all()
    return render_template("upgrade.html", my_requests=my_requests)


# ── ADMIN: generate license codes ────────────────────────────────────────────
@app.route("/admin/licenses", methods=["GET", "POST"])
@login_required
def admin_licenses():
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        plan = request.form.get("plan", "pro")
        days = int(request.form.get("days", 365))
        lic = LicenseCode(plan=plan, duration_days=days)
        db.session.add(lic)
        log_admin_action("license_generated", target=f"plan={plan}", details=f"{days} days")
        db.session.commit()
        flash(tr("flash_license_generated", code=lic.code), "success")

    q = request.args.get("q", "").strip()
    status_filter = request.args.get("status", "").strip()
    page = request.args.get("page", 1, type=int)

    query = LicenseCode.query
    if q:
        query = query.filter(LicenseCode.code.ilike(f"%{q}%"))
    if status_filter == "used":
        query = query.filter(LicenseCode.used.is_(True))
    elif status_filter == "available":
        query = query.filter(LicenseCode.used.is_(False))

    query = query.order_by(LicenseCode.created_at.desc())
    pager = query.paginate(page=page, per_page=20, error_out=False)

    return render_template("admin_licenses.html", codes=pager.items, pager=pager,
                            q=q, status_filter=status_filter)


# ── ADMIN: payment validation — activate / renew pro & premium subscriptions ─
@app.route("/admin/payments")
@login_required
def admin_payments():
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    status_filter = request.args.get("status", "pending").strip()
    q = request.args.get("q", "").strip()
    page = request.args.get("page", 1, type=int)

    query = PaymentRequest.query.join(User, PaymentRequest.user_id == User.id)
    if status_filter in ("pending", "approved", "rejected"):
        query = query.filter(PaymentRequest.status == status_filter)
    if q:
        like = f"%{q}%"
        query = query.filter(db.or_(User.first_name.ilike(like), User.last_name.ilike(like),
                                     User.email.ilike(like), PaymentRequest.reference.ilike(like)))

    query = query.order_by(PaymentRequest.created_at.desc())
    pager = query.paginate(page=page, per_page=20, error_out=False)

    return render_template("admin_payments.html", pager=pager, q=q, status_filter=status_filter)


@app.route("/admin/payments/<int:req_id>/approve", methods=["POST"])
@login_required
def admin_approve_payment(req_id):
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    preq = PaymentRequest.query.get_or_404(req_id)
    if preq.status != "pending":
        flash(tr("flash_payment_already_reviewed"), "error")
        return redirect(url_for("admin_payments"))

    days = request.form.get("days", 365, type=int)
    is_renewal = apply_payment_approval(preq.user, preq.plan, days=days)

    log_payment(preq.user_id, preq.plan, amount=preq.amount_claimed,
                note=f"PaymentRequest #{preq.id} ({preq.reference})")

    preq.status = "approved"
    preq.reviewed_at = datetime.utcnow()
    preq.reviewed_by = current_user.id

    kind = "renewal" if is_renewal else "activation"
    log_admin_action("payment_approved", target=preq.user.email,
                      details=f"plan={preq.plan} {kind} +{days}d")

    notif_body = tr("msg_payment_approved", plan=preq.user.subscription.plan_label(),
                     date=preq.user.subscription.expires_at.strftime("%d/%m/%Y"))
    db.session.add(Message(user_id=preq.user_id, sender="admin", body=notif_body,
                            read_by_admin=True, read_by_teacher=False))

    db.session.commit()
    flash(tr("flash_payment_approved", name=preq.user.full_name), "success")
    return redirect(url_for("admin_payments"))


@app.route("/admin/payments/<int:req_id>/reject", methods=["POST"])
@login_required
def admin_reject_payment(req_id):
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    preq = PaymentRequest.query.get_or_404(req_id)
    if preq.status != "pending":
        flash(tr("flash_payment_already_reviewed"), "error")
        return redirect(url_for("admin_payments"))

    reason = request.form.get("reason", "").strip()
    preq.status = "rejected"
    preq.reviewed_at = datetime.utcnow()
    preq.reviewed_by = current_user.id
    preq.admin_note = reason

    log_admin_action("payment_rejected", target=preq.user.email, details=reason)

    notif_body = tr("msg_payment_rejected", reason=reason or tr("msg_payment_rejected_generic"))
    db.session.add(Message(user_id=preq.user_id, sender="admin", body=notif_body,
                            read_by_admin=True, read_by_teacher=False))

    db.session.commit()
    flash(tr("flash_payment_rejected", name=preq.user.full_name), "success")
    return redirect(url_for("admin_payments"))


# ── ADMIN: advanced statistics dashboard ─────────────────────────────────────
@app.route("/admin/stats")
@login_required
def admin_stats():
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    from sqlalchemy import func

    total_teachers = User.query.filter_by(role="user").count()
    plan_counts = {}
    for p in ("trial", "pro", "premium"):
        plan_counts[p] = Subscription.query.filter_by(plan=p).count()

    total_docs = Document.query.count()

    # Documents per day, last 14 days
    since = datetime.utcnow() - timedelta(days=14)
    daily_rows = (
        db.session.query(func.date(Document.created_at), func.count(Document.id))
        .filter(Document.created_at >= since)
        .group_by(func.date(Document.created_at))
        .all()
    )
    daily_map = {str(d): c for d, c in daily_rows}
    chart_labels, chart_values = [], []
    for i in range(13, -1, -1):
        day = (datetime.utcnow() - timedelta(days=i)).date()
        chart_labels.append(day.strftime("%d/%m"))
        chart_values.append(daily_map.get(str(day), 0))

    # Revenue
    total_revenue = db.session.query(func.sum(Payment.amount)).scalar() or 0
    month_start = datetime.utcnow().replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    month_revenue = (
        db.session.query(func.sum(Payment.amount))
        .filter(Payment.created_at >= month_start)
        .scalar() or 0
    )
    recent_payments = Payment.query.order_by(Payment.created_at.desc()).limit(20).all()

    # Expiring within 5 days
    now = datetime.utcnow()
    soon = now + timedelta(days=5)
    expiring = (
        Subscription.query.filter(
            Subscription.status == "active",
            Subscription.expires_at.isnot(None),
            Subscription.expires_at >= now,
            Subscription.expires_at <= soon,
        ).all()
    )

    # Top 5 most active teachers by number of documents generated
    top_rows = (
        db.session.query(User, func.count(Document.id).label("cnt"))
        .join(Document, Document.user_id == User.id)
        .group_by(User.id)
        .order_by(func.count(Document.id).desc())
        .limit(5)
        .all()
    )

    pending_payments = PaymentRequest.query.filter_by(status="pending").count()

    return render_template("admin_stats.html",
                           total_teachers=total_teachers, plan_counts=plan_counts,
                           total_docs=total_docs, chart_labels=chart_labels, chart_values=chart_values,
                           total_revenue=total_revenue, month_revenue=month_revenue,
                           recent_payments=recent_payments, expiring=expiring,
                           top_teachers=top_rows, pending_payments=pending_payments)


# ── ADMIN: audit log (accountability trail of admin actions) ────────────────
@app.route("/admin/audit-log")
@login_required
def admin_audit_log():
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    q = request.args.get("q", "").strip()
    page = request.args.get("page", 1, type=int)

    query = AdminLog.query
    if q:
        like = f"%{q}%"
        query = query.filter(db.or_(AdminLog.action.ilike(like), AdminLog.target.ilike(like),
                                     AdminLog.details.ilike(like)))
    query = query.order_by(AdminLog.created_at.desc())
    pager = query.paginate(page=page, per_page=30, error_out=False)

    return render_template("admin_audit_log.html", pager=pager, q=q)


def generate_password(length=12):
    alphabet = string.ascii_letters + string.digits
    return ''.join(secrets.choice(alphabet) for _ in range(length))


# ── ADMIN: create & manage teacher accounts ──────────────────────────────────
@app.route("/admin/teachers", methods=["GET", "POST"])
@login_required
def admin_teachers():
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    generated = None  # (email, password) shown once after creation

    if request.method == "POST":
        first_name = request.form.get("first_name", "").strip()
        last_name  = request.form.get("last_name", "").strip()
        email      = request.form.get("email", "").strip().lower()
        plan       = request.form.get("plan", "trial")
        custom_pw  = request.form.get("password", "").strip()
        api_key    = request.form.get("api_key", "").strip()

        if not all([first_name, last_name, email]):
            flash(tr("flash_teacher_fields_required"), "error")
        elif User.query.filter_by(email=email).first():
            flash(tr("flash_email_exists"), "error")
        else:
            password = custom_pw or generate_password()
            user = User(first_name=first_name, last_name=last_name, email=email,
                        gemini_api_key=api_key or None)
            user.set_password(password)
            db.session.add(user)
            db.session.flush()

            if plan == "trial":
                create_trial_subscription(user)
            else:
                db.session.add(Subscription(
                    user_id=user.id, plan=plan, status="active",
                    expires_at=datetime.utcnow() + timedelta(days=365),
                    docs_used=0, docs_used_today=0, last_reset_date=None,
                ))
                log_payment(user.id, plan, note="Created by admin")

            log_admin_action("teacher_created", target=email, details=f"plan={plan}")
            db.session.commit()
            generated = (email, password)
            flash(tr("flash_teacher_created", name=f"{first_name} {last_name}"), "success")

    q = request.args.get("q", "").strip()
    plan_filter = request.args.get("plan", "").strip()
    page = request.args.get("page", 1, type=int)

    query = User.query.filter_by(role="user")
    if q:
        like = f"%{q}%"
        query = query.filter(db.or_(User.first_name.ilike(like),
                                     User.last_name.ilike(like),
                                     User.email.ilike(like)))
    if plan_filter:
        query = query.join(Subscription).filter(Subscription.plan == plan_filter)

    query = query.order_by(User.created_at.desc())
    pager = query.paginate(page=page, per_page=15, error_out=False)

    return render_template("admin_teachers.html", teachers=pager.items, pager=pager,
                            q=q, plan_filter=plan_filter, generated=generated)


@app.route("/admin/teachers/export.csv")
@login_required
def admin_teachers_export():
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    import csv, io as _io
    q = request.args.get("q", "").strip()
    plan_filter = request.args.get("plan", "").strip()

    query = User.query.filter_by(role="user")
    if q:
        like = f"%{q}%"
        query = query.filter(db.or_(User.first_name.ilike(like),
                                     User.last_name.ilike(like),
                                     User.email.ilike(like)))
    if plan_filter:
        query = query.join(Subscription).filter(Subscription.plan == plan_filter)
    teachers = query.order_by(User.created_at.desc()).all()

    buf = _io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["#", "Prénom", "Nom", "Email", "Abonnement", "Statut", "Expire le",
                      "Documents générés", "Créé le", "Actif"])
    for i, t in enumerate(teachers, start=1):
        sub = t.subscription
        writer.writerow([
            i, t.first_name, t.last_name, t.email,
            sub.plan if sub else "", sub.status if sub else "",
            sub.expires_at.strftime("%d/%m/%Y") if sub and sub.expires_at else "",
            len(t.documents), t.created_at.strftime("%d/%m/%Y"),
            "Oui" if t.is_active else "Non",
        ])
    log_admin_action("export_csv", target="teachers")
    db.session.commit()
    return app.response_class(buf.getvalue(), mimetype="text/csv",
                               headers={"Content-Disposition": "attachment; filename=enseignants.csv"})


@app.route("/admin/teachers/<int:user_id>")
@login_required
def admin_teacher_detail(user_id):
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    teacher = User.query.get_or_404(user_id)
    sub = teacher.subscription

    q = request.args.get("q", "").strip()
    page = request.args.get("page", 1, type=int)

    doc_query = Document.query.filter_by(user_id=teacher.id)
    if q:
        like = f"%{q}%"
        doc_query = doc_query.filter(db.or_(Document.title.ilike(like),
                                             Document.subject.ilike(like),
                                             Document.doc_type.ilike(like),
                                             Document.level.ilike(like)))
    doc_query = doc_query.order_by(Document.created_at.desc())
    docs_pager = doc_query.paginate(page=page, per_page=10, error_out=False)

    payments = Payment.query.filter_by(user_id=teacher.id).order_by(Payment.created_at.desc()).all()
    total_docs = Document.query.filter_by(user_id=teacher.id).count()

    return render_template("admin_teacher_detail.html", teacher=teacher, sub=sub,
                            docs_pager=docs_pager, payments=payments, q=q,
                            total_docs=total_docs)


@app.route("/admin/teachers/<int:user_id>/reset-quota", methods=["POST"])
@login_required
def admin_reset_quota(user_id):
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    teacher = User.query.get_or_404(user_id)
    sub = teacher.subscription
    if sub:
        sub.docs_used_today = 0
        sub.docs_used = 0
        sub.last_reset_date = None
        log_admin_action("quota_reset", target=teacher.email)
        db.session.commit()
        flash(tr("flash_quota_reset", name=teacher.full_name), "success")
    return redirect(url_for("admin_teacher_detail", user_id=user_id))


@app.route("/admin/teachers/<int:user_id>/extend", methods=["POST"])
@login_required
def admin_extend_subscription(user_id):
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    teacher = User.query.get_or_404(user_id)
    sub = teacher.subscription
    days = request.form.get("days", 30, type=int)
    if sub:
        base = sub.expires_at if (sub.expires_at and sub.expires_at > datetime.utcnow()) else datetime.utcnow()
        sub.expires_at = base + timedelta(days=days)
        sub.status = "active"
        sub.expiry_reminder_sent = False
        log_admin_action("subscription_extended", target=teacher.email, details=f"+{days}d")
        db.session.commit()
        flash(tr("flash_subscription_extended", name=teacher.full_name, days=days), "success")
    return redirect(url_for("admin_teacher_detail", user_id=user_id))


# ── ADMIN: cross-teacher document search ─────────────────────────────────────
@app.route("/admin/documents")
@login_required
def admin_documents():
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    q = request.args.get("q", "").strip()
    doc_type = request.args.get("doc_type", "").strip()
    page = request.args.get("page", 1, type=int)

    query = Document.query.join(User, Document.user_id == User.id)
    if q:
        like = f"%{q}%"
        query = query.filter(db.or_(Document.title.ilike(like),
                                     Document.subject.ilike(like),
                                     Document.level.ilike(like),
                                     User.first_name.ilike(like),
                                     User.last_name.ilike(like),
                                     User.email.ilike(like)))
    if doc_type:
        query = query.filter(Document.doc_type == doc_type)

    query = query.order_by(Document.created_at.desc())
    pager = query.paginate(page=page, per_page=20, error_out=False)

    doc_type_options = [row[0] for row in
                         db.session.query(Document.doc_type).distinct().all() if row[0]]

    return render_template("admin_documents.html", pager=pager, q=q,
                            doc_type=doc_type, doc_type_options=doc_type_options)


@app.route("/admin/documents/export.csv")
@login_required
def admin_documents_export():
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    import csv, io as _io
    q = request.args.get("q", "").strip()
    doc_type = request.args.get("doc_type", "").strip()

    query = Document.query.join(User, Document.user_id == User.id)
    if q:
        like = f"%{q}%"
        query = query.filter(db.or_(Document.title.ilike(like), Document.subject.ilike(like),
                                     Document.level.ilike(like), User.first_name.ilike(like),
                                     User.last_name.ilike(like), User.email.ilike(like)))
    if doc_type:
        query = query.filter(Document.doc_type == doc_type)
    docs = query.order_by(Document.created_at.desc()).all()

    buf = _io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["#", "Enseignant", "Email", "Titre", "Matière", "Niveau", "Type",
                      "Format", "Date"])
    for i, d in enumerate(docs, start=1):
        writer.writerow([i, d.user.full_name if d.user else "", d.user.email if d.user else "",
                          d.title, d.subject, d.level, d.doc_type, d.fmt,
                          d.created_at.strftime("%d/%m/%Y %H:%M")])
    log_admin_action("export_csv", target="documents")
    db.session.commit()
    return app.response_class(buf.getvalue(), mimetype="text/csv",
                               headers={"Content-Disposition": "attachment; filename=documents.csv"})


# ── MESSAGING: teacher <-> admin, free & unlimited ───────────────────────────
@app.route("/messages", methods=["GET", "POST"])
@login_required
def messages():
    if current_user.role == "admin":
        return redirect(url_for("admin_messages"))

    if request.method == "POST":
        body = request.form.get("body", "").strip()
        if body:
            db.session.add(Message(user_id=current_user.id, sender="teacher", body=body,
                                    read_by_teacher=True, read_by_admin=False))
            db.session.commit()
        return redirect(url_for("messages"))

    # Mark admin's replies as read now that the teacher is viewing the thread
    Message.query.filter_by(user_id=current_user.id, sender="admin", read_by_teacher=False) \
        .update({"read_by_teacher": True})
    db.session.commit()

    thread = Message.query.filter_by(user_id=current_user.id).order_by(Message.created_at.asc()).all()
    return render_template("messages.html", thread=thread)


@app.route("/admin/messages")
@login_required
def admin_messages():
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    q = request.args.get("q", "").strip()
    page = request.args.get("page", 1, type=int)

    teachers_query = User.query.filter_by(role="user")
    if q:
        like = f"%{q}%"
        teachers_query = teachers_query.filter(db.or_(User.first_name.ilike(like),
                                                        User.last_name.ilike(like),
                                                        User.email.ilike(like)))
    teachers = teachers_query.all()

    threads = []
    for t in teachers:
        last = Message.query.filter_by(user_id=t.id).order_by(Message.created_at.desc()).first()
        unread = Message.query.filter_by(user_id=t.id, sender="teacher", read_by_admin=False).count()
        if last:  # only list teachers who have exchanged at least one message
            threads.append({"teacher": t, "last": last, "unread": unread})

    threads.sort(key=lambda x: (x["unread"] == 0, -x["last"].created_at.timestamp()))

    page_items, page, total_pages, total = paginate_list(threads, page, per_page=15)

    return render_template("admin_messages.html", threads=page_items, q=q,
                            page=page, total_pages=total_pages, total=total)


@app.route("/admin/messages/<int:user_id>", methods=["GET", "POST"])
@login_required
def admin_message_thread(user_id):
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    teacher = User.query.get_or_404(user_id)

    if request.method == "POST":
        body = request.form.get("body", "").strip()
        if body:
            db.session.add(Message(user_id=teacher.id, sender="admin", body=body,
                                    read_by_admin=True, read_by_teacher=False))
            db.session.commit()
        return redirect(url_for("admin_message_thread", user_id=user_id))

    Message.query.filter_by(user_id=teacher.id, sender="teacher", read_by_admin=False) \
        .update({"read_by_admin": True})
    db.session.commit()

    thread = Message.query.filter_by(user_id=teacher.id).order_by(Message.created_at.asc()).all()
    return render_template("admin_message_thread.html", teacher=teacher, thread=thread)


@app.route("/admin/teachers/<int:user_id>/edit", methods=["GET", "POST"])
@login_required
def admin_edit_teacher(user_id):
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    teacher = User.query.get_or_404(user_id)
    sub = teacher.subscription

    if request.method == "POST":
        first_name = request.form.get("first_name", "").strip()
        last_name  = request.form.get("last_name", "").strip()
        email      = request.form.get("email", "").strip().lower()
        plan       = request.form.get("plan", sub.plan if sub else "trial")
        api_key    = request.form.get("api_key", "").strip()

        existing = User.query.filter_by(email=email).first()
        if existing and existing.id != teacher.id:
            flash(tr("flash_email_exists"), "error")
            return render_template("admin_edit_teacher.html", teacher=teacher, sub=sub)

        teacher.first_name = first_name
        teacher.last_name = last_name
        teacher.email = email
        teacher.gemini_api_key = api_key or teacher.gemini_api_key

        if sub and plan != sub.plan:
            old_plan = sub.plan
            sub.plan = plan
            sub.docs_used_today = 0
            sub.last_reset_date = None
            sub.expiry_reminder_sent = False
            if plan == "trial" and old_plan != "trial":
                sub.docs_used = 0
                sub.docs_limit = 5
                sub.expires_at = datetime.utcnow() + timedelta(days=14)
            elif plan != "trial":
                sub.expires_at = datetime.utcnow() + timedelta(days=365)
                log_payment(teacher.id, plan, note="Plan changed by admin")

        log_admin_action("teacher_updated", target=teacher.email)
        db.session.commit()
        flash(tr("flash_teacher_updated", name=teacher.full_name), "success")
        return redirect(url_for("admin_teachers"))

    return render_template("admin_edit_teacher.html", teacher=teacher, sub=sub)


@app.route("/admin/teachers/<int:user_id>/reset-password", methods=["POST"])
@login_required
def admin_reset_password(user_id):
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    user = User.query.get_or_404(user_id)
    new_password = generate_password()
    user.set_password(new_password)
    log_admin_action("password_reset", target=user.email)
    db.session.commit()
    flash(tr("flash_password_reset", email=user.email, password=new_password), "success")
    return redirect(url_for("admin_teachers"))


@app.route("/admin/teachers/<int:user_id>/delete", methods=["POST"])
@login_required
def admin_delete_teacher(user_id):
    """Soft-delete: deactivates the account (blocks login) rather than erasing it,
    so document history, payments and messages are preserved. Toggles back on if
    the account is already deactivated (reactivation)."""
    if current_user.role != "admin":
        flash(tr("flash_admin_only"), "error")
        return redirect(url_for("dashboard"))

    user = User.query.get_or_404(user_id)
    redirect_to = request.form.get("redirect_to") or url_for("admin_teachers")

    if user.role == "admin":
        flash(tr("flash_admin_delete_protected"), "error")
    elif user.is_active:
        user.is_active = False
        log_admin_action("teacher_deactivated", target=user.email)
        db.session.commit()
        flash(tr("flash_teacher_deactivated", email=user.email), "success")
    else:
        user.is_active = True
        log_admin_action("teacher_reactivated", target=user.email)
        db.session.commit()
        flash(tr("flash_teacher_reactivated", email=user.email), "success")
    return redirect(redirect_to)


# ── MAIN GENERATOR PAGE ───────────────────────────────────────────────────────
@app.route("/generator")
@login_required
def generator():
    return render_template("generator.html",
                           levels=LEVELS, subjects=SUBJECTS, doc_types=DOCUMENT_TYPES)


# ── API: QUOTA CHECK ──────────────────────────────────────────────────────────
@app.route("/api/check-quota")
@login_required
def check_quota():
    sub = current_user.subscription
    allowed = sub.has_quota()
    db.session.commit()  # persist any daily-counter roll-over from has_quota()
    return jsonify({
        "allowed": allowed,
        "remaining": sub.remaining(),
        "plan": sub.plan,
        "plan_label": sub.plan_label(),
        "is_daily": sub.is_daily_plan(),
        "daily_limit": sub.daily_limit() if sub.is_daily_plan() else None,
        "contact_email": CONTACT_EMAIL,
    })


# ── API: LOG USAGE (after successful AI generation) ──────────────────────────
@app.route("/api/log-usage", methods=["POST"])
@login_required
@csrf.exempt
@limiter.limit("60/hour")
def log_usage():
    data = request.get_json() or {}
    sub = current_user.subscription

    if not sub.has_quota():
        return jsonify({"error": "quota_exceeded", "contact_email": CONTACT_EMAIL}), 403

    sub.register_usage()

    doc = Document(
        user_id=current_user.id,
        title=data.get("lesson", "")[:255],
        subject=data.get("subject", ""),
        level=data.get("level", ""),
        palier=data.get("palier", ""),
        doc_type=data.get("doc_type", ""),
        fmt=data.get("format", "docx"),
    )
    db.session.add(doc)
    db.session.commit()

    return jsonify({"ok": True, "remaining": sub.remaining()})


# ── API: EXTRACT SOURCES (NotebookLM-style: files + links) ──────────────────
@app.route("/api/extract-sources", methods=["POST"])
@login_required
@csrf.exempt
@limiter.limit("30/hour")
def extract_sources():
    files = request.files.getlist("files")
    urls_raw = request.form.get("urls", "")
    urls = [u.strip() for u in urls_raw.split("\n") if u.strip()]

    context, warnings = build_sources_context(files, urls)
    return jsonify({"context": context, "warnings": warnings})


# ── DOCUMENT GENERATION HELPERS ───────────────────────────────────────────────
def hex_to_rgb(hex_color, default="102a53"):
    h = (hex_color or default).lstrip("#")
    if len(h) != 6:
        h = default
    return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))


def set_cell_background(cell, hex_color):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), hex_color.lstrip("#"))
    cell._tc.get_or_add_tcPr().append(shd)


def try_parse_md_table(lines, i):
    """If lines[i] starts a Markdown pipe-table, parse it and return (rows, next_index).
    Otherwise return None. Handles the |---|---| separator row Gemini emits."""
    def is_row(s):
        s = s.strip()
        return s.startswith("|") and s.endswith("|") and s.count("|") >= 2

    def is_sep(s):
        s = s.strip().strip("|")
        return bool(s) and all(c in "-:| " for c in s) and "-" in s

    if i + 1 >= len(lines) or not is_row(lines[i]) or not is_sep(lines[i + 1]):
        return None

    def split_row(s):
        s = s.strip()
        if s.startswith("|"):
            s = s[1:]
        if s.endswith("|"):
            s = s[:-1]
        return [c.strip() for c in s.split("|")]

    rows = [split_row(lines[i])]
    j = i + 2
    while j < len(lines) and is_row(lines[j]):
        rows.append(split_row(lines[j]))
        j += 1
    return rows, j


HEADER_LABELS = {
    "fr": {"teacher": "Enseignant(e)", "etablissement": "Établissement", "level": "Niveau",
           "palier": "Palier", "subject": "Matière", "projet": "Projet / Unité", "activite": "Activité"},
    "en": {"teacher": "Teacher", "etablissement": "Institution", "level": "Level",
           "palier": "Grade", "subject": "Subject", "projet": "Project / Unit", "activite": "Activity"},
    "ar": {"teacher": "الأستاذ(ة)", "etablissement": "المؤسسة", "level": "المستوى",
           "palier": "الطور", "subject": "المادة", "projet": "المشروع / الوحدة", "activite": "النشاط"},
}


def make_docx(content, meta, colors_cfg, lang="fr"):
    """
    meta: dict with teacher_first, teacher_last, level, palier, subject, lesson
    colors_cfg: dict with c1 (header bg), c2 (H1), c3 (accent/H2), text (body text)
    """
    c1 = colors_cfg.get("c1", "#102a53")
    c2 = colors_cfg.get("c2", "#f97316")
    c3 = colors_cfg.get("c3", "#173f6c")
    ctext = colors_cfg.get("text", "#1a1a1a")

    doc = DocxDocument()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ── Brand title ──
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run("HaithemEduAI")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(*hex_to_rgb(c1))
    title_p.paragraph_format.space_after = Pt(14)

    # ── Header table: vertical label:value rows (scales to any number of fields) ──
    labels = HEADER_LABELS.get(lang, HEADER_LABELS["fr"])
    rows_data = [
        (labels["teacher"], f"{meta.get('teacher_first','')} {meta.get('teacher_last','')}".strip() or "—"),
        (labels["etablissement"], meta.get("etablissement") or "—"),
        (labels["level"], meta.get("level") or "—"),
        (labels["palier"], meta.get("palier") or "—"),
        (labels["subject"], meta.get("subject") or "—"),
        (labels["projet"], meta.get("projet") or "—"),
        (labels["activite"], meta.get("activite") or "—"),
    ]

    table = doc.add_table(rows=len(rows_data), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(4.2)

    for i, (h, v) in enumerate(rows_data):
        hc = table.cell(i, 0)
        hc.text = ""
        hp = hc.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT if lang != "ar" else WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run(h)
        hr.bold = True
        hr.font.size = Pt(10)
        hr.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hc, c1)
        hc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        vc = table.cell(i, 1)
        vc.text = ""
        vp = vc.paragraphs[0]
        vp.alignment = WD_ALIGN_PARAGRAPH.LEFT if lang != "ar" else WD_ALIGN_PARAGRAPH.RIGHT
        vr = vp.add_run(v)
        vr.font.size = Pt(11)
        vr.font.color.rgb = RGBColor(*hex_to_rgb(ctext, "1a1a1a"))
        vc.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ── Lesson topic banner ──
    lesson_p = doc.add_paragraph()
    lesson_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lr = lesson_p.add_run(meta.get("lesson", ""))
    lr.bold = True
    lr.font.size = Pt(14)
    lr.font.color.rgb = RGBColor(*hex_to_rgb(c3))
    lesson_p.paragraph_format.space_after = Pt(16)

    # ── Body content (markdown-ish parsing, with real tables) ──
    lines = content.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].strip()

        if not line:
            doc.add_paragraph("")
            i += 1
            continue

        table_result = try_parse_md_table(lines, i)
        if table_result:
            rows, i = table_result
            n_cols = max(len(r) for r in rows)
            tbl = doc.add_table(rows=len(rows), cols=n_cols)
            tbl.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci in range(n_cols):
                    cell = tbl.cell(ri, ci)
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if lang == "ar" else WD_ALIGN_PARAGRAPH.LEFT
                    r = p.add_run(row[ci] if ci < len(row) else "")
                    r.font.size = Pt(10)
                    if ri == 0:
                        r.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        set_cell_background(cell, c3)
                    else:
                        r.font.color.rgb = RGBColor(*hex_to_rgb(ctext, "1a1a1a"))
            doc.add_paragraph().paragraph_format.space_after = Pt(10)
            continue

        if line.startswith("### "):
            h = doc.add_heading("", level=3)
            r = h.add_run(line[4:])
            r.font.color.rgb = RGBColor(*hex_to_rgb(c3))
        elif line.startswith("## "):
            h = doc.add_heading("", level=2)
            r = h.add_run(line[3:])
            r.font.color.rgb = RGBColor(*hex_to_rgb(c3))
        elif line.startswith("# "):
            h = doc.add_heading("", level=1)
            r = h.add_run(line[2:])
            r.font.color.rgb = RGBColor(*hex_to_rgb(c2))
        elif line.startswith("- ") or line.startswith("• "):
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(line[2:])
            r.font.color.rgb = RGBColor(*hex_to_rgb(ctext, "1a1a1a"))
        else:
            p = doc.add_paragraph()
            parts = line.split("**")
            for pi, part in enumerate(parts):
                r = p.add_run(part)
                r.font.color.rgb = RGBColor(*hex_to_rgb(ctext, "1a1a1a"))
                if pi % 2 == 1:
                    r.bold = True

        i += 1

    # ── Footer ──
    doc.add_paragraph().paragraph_format.space_before = Pt(20)
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer_p.add_run("Généré avec HaithemEduAI")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(150, 150, 150)
    fr.italic = True

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def make_pdf(content, meta, colors_cfg, lang="fr"):
    c1 = colors_cfg.get("c1", "#102a53")
    c2 = colors_cfg.get("c2", "#f97316")
    c3 = colors_cfg.get("c3", "#173f6c")
    ctext = colors_cfg.get("text", "#1a1a1a")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            rightMargin=2*cm, leftMargin=2*cm,
                            topMargin=1.6*cm, bottomMargin=1.6*cm)
    styles = getSampleStyleSheet()
    align = TA_RIGHT if lang == "ar" else TA_LEFT

    title_s = ParagraphStyle("T", parent=styles["Title"], textColor=rl_colors.HexColor(c1),
                              fontSize=20, spaceAfter=10, alignment=TA_CENTER)
    lesson_s = ParagraphStyle("L", parent=styles["Normal"], textColor=rl_colors.HexColor(c3),
                               fontSize=14, spaceAfter=14, alignment=TA_CENTER, fontName="Helvetica-Bold")
    h1_s = ParagraphStyle("H1", parent=styles["Heading1"], textColor=rl_colors.HexColor(c2),
                           fontSize=14, spaceBefore=14, spaceAfter=4, alignment=align)
    h2_s = ParagraphStyle("H2", parent=styles["Heading2"], textColor=rl_colors.HexColor(c3),
                           fontSize=12, spaceBefore=10, spaceAfter=3, alignment=align)
    body_s = ParagraphStyle("B", parent=styles["Normal"], textColor=rl_colors.HexColor(ctext),
                             fontSize=10, leading=16, spaceAfter=5, alignment=align)
    blt_s = ParagraphStyle("Bl", parent=styles["Normal"], textColor=rl_colors.HexColor(ctext),
                            fontSize=10, leading=16, leftIndent=20, spaceAfter=3, alignment=align)

    story = [Paragraph("HaithemEduAI", title_s)]

    teacher = f"{meta.get('teacher_first','')} {meta.get('teacher_last','')}".strip() or "—"
    labels = HEADER_LABELS.get(lang, HEADER_LABELS["fr"])
    table_data = [
        [labels["teacher"], teacher],
        [labels["etablissement"], meta.get("etablissement") or "—"],
        [labels["level"], meta.get("level") or "—"],
        [labels["palier"], meta.get("palier") or "—"],
        [labels["subject"], meta.get("subject") or "—"],
        [labels["projet"], meta.get("projet") or "—"],
        [labels["activite"], meta.get("activite") or "—"],
    ]
    if lang == "ar":
        table_data = [[row[1], row[0]] for row in table_data]  # value first visually for RTL

    tbl = Table(table_data, colWidths=[5*cm, 9*cm])
    tbl.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (0,-1) if lang != "ar" else (1,-1), rl_colors.HexColor(c1)),
        ("TEXTCOLOR", (0,0), (0,-1) if lang != "ar" else (1,-1), rl_colors.white),
        ("TEXTCOLOR", (1,0), (1,-1) if lang != "ar" else (0,-1), rl_colors.HexColor(ctext)),
        ("FONTNAME", (0,0), (0,-1) if lang != "ar" else (1,-1), "Helvetica-Bold"),
        ("FONTSIZE", (0,0), (-1,-1), 9),
        ("ALIGN", (0,0), (-1,-1), "RIGHT" if lang == "ar" else "LEFT"),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("GRID", (0,0), (-1,-1), 0.75, rl_colors.HexColor("#dce6f0")),
        ("TOPPADDING", (0,0), (-1,-1), 7),
        ("BOTTOMPADDING", (0,0), (-1,-1), 7),
        ("LEFTPADDING", (0,0), (-1,-1), 10),
    ]))
    story.append(tbl)
    story.append(Spacer(1, 0.5*cm))
    story.append(Paragraph(meta.get("lesson",""), lesson_s))

    cell_s = ParagraphStyle("Cell", parent=styles["Normal"], textColor=rl_colors.HexColor(ctext),
                             fontSize=9, leading=13, alignment=align)
    cell_head_s = ParagraphStyle("CellH", parent=styles["Normal"], textColor=rl_colors.white,
                                  fontSize=9, leading=13, alignment=align, fontName="Helvetica-Bold")

    lines = content.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].strip()

        if not line:
            story.append(Spacer(1, 0.2*cm))
            i += 1
            continue

        table_result = try_parse_md_table(lines, i)
        if table_result:
            rows, i = table_result
            n_cols = max(len(r) for r in rows)
            wrapped = []
            for ri, row in enumerate(rows):
                style = cell_head_s if ri == 0 else cell_s
                wrapped.append([Paragraph(row[ci] if ci < len(row) else "", style) for ci in range(n_cols)])
            col_w = (16*cm) / n_cols
            body_tbl = Table(wrapped, colWidths=[col_w]*n_cols)
            body_tbl.setStyle(TableStyle([
                ("BACKGROUND", (0,0), (-1,0), rl_colors.HexColor(c3)),
                ("GRID", (0,0), (-1,-1), 0.6, rl_colors.HexColor("#dce6f0")),
                ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
                ("TOPPADDING", (0,0), (-1,-1), 5),
                ("BOTTOMPADDING", (0,0), (-1,-1), 5),
            ]))
            story.append(body_tbl)
            story.append(Spacer(1, 0.3*cm))
            continue

        if line.startswith("### ") or line.startswith("## ") or line.startswith("# "):
            style = h1_s if line.startswith("# ") else h2_s
            story.append(Paragraph(line.lstrip("#").strip(), style))
        elif line.startswith("- ") or line.startswith("• "):
            story.append(Paragraph("• " + line[2:], blt_s))
        else:
            safe = line.replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")
            safe = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', safe)
            story.append(Paragraph(safe, body_s))

        i += 1

    doc.build(story)
    buf.seek(0)
    return buf


# ── API: DOWNLOAD ──────────────────────────────────────────────────────────────
@app.route("/api/download", methods=["POST"])
@login_required
@csrf.exempt
@limiter.limit("60/hour")
def download():
    data = request.get_json()
    content = data.get("content", "")
    lang    = data.get("lang", "fr")
    fmt     = data.get("format", "docx")

    meta = {
        "teacher_first": data.get("teacher_first", current_user.first_name),
        "teacher_last":  data.get("teacher_last", current_user.last_name),
        "level":         data.get("level", ""),
        "palier":        data.get("palier", ""),
        "subject":       data.get("subject", ""),
        "lesson":        data.get("lesson", ""),
        "etablissement": data.get("etablissement", ""),
        "projet":        data.get("projet", ""),
        "activite":      data.get("activite", ""),
    }
    colors_cfg = data.get("colors", {})

    safe_name = re.sub(r'[^\w\-]', '_', meta["lesson"][:40] or "document")

    if fmt == "pdf":
        buf = make_pdf(content, meta, colors_cfg, lang)
        return send_file(buf, as_attachment=True, download_name=f"{safe_name}.pdf",
                         mimetype="application/pdf")
    else:
        buf = make_docx(content, meta, colors_cfg, lang)
        return send_file(buf, as_attachment=True, download_name=f"{safe_name}.docx",
                         mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


if __name__ == "__main__":
    app.run(debug=True, port=5000)
